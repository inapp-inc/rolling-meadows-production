"""Translation catalogue helpers shared by platform and tenant admin routes."""

from __future__ import annotations

import json
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.i18n import TranslationEntry, TranslationLocale

SEED_EN_JSON = Path(__file__).resolve().parents[2] / "seed-data" / "en.json"
FALLBACK_EN_JSON = Path(__file__).resolve().parents[3] / "web" / "src" / "i18n" / "locales" / "en.json"
LOCALES_DIR = Path(__file__).resolve().parents[3] / "web" / "src" / "i18n" / "locales"


def flatten_json(obj: dict[str, Any], prefix: str = "") -> dict[str, str]:
    result: dict[str, str] = {}
    for key, value in obj.items():
        full_key = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            result.update(flatten_json(value, full_key))
        elif isinstance(value, str):
            result[full_key] = value
    return result


def load_static_locale_bundle(locale: str) -> dict[str, str]:
    path = LOCALES_DIR / f"{locale.strip().lower()}.json"
    if not path.is_file():
        return {}
    with path.open(encoding="utf-8") as handle:
        return flatten_json(json.load(handle))


PROTECTED_LOCALE_CODES = frozenset({"en"})


async def list_active_locales(session: AsyncSession) -> list[TranslationLocale]:
    result = await session.execute(
        select(TranslationLocale)
        .where(TranslationLocale.active.is_(True))
        .order_by(TranslationLocale.code)
    )
    return list(result.scalars().all())


async def list_all_locales(session: AsyncSession) -> list[TranslationLocale]:
    result = await session.execute(select(TranslationLocale).order_by(TranslationLocale.code))
    return list(result.scalars().all())


async def load_master_keys(session: AsyncSession) -> dict[str, str]:
    result = await session.execute(
        select(TranslationEntry).where(
            TranslationEntry.tenant_id.is_(None),
            TranslationEntry.locale == "en",
        )
    )
    entries = result.scalars().all()
    if entries:
        return {entry.key: entry.value for entry in entries}

    for en_path in (SEED_EN_JSON, FALLBACK_EN_JSON):
        if en_path.is_file():
            with en_path.open(encoding="utf-8") as handle:
                return flatten_json(json.load(handle))
    return {}


async def _load_db_values(
    session: AsyncSession,
    *,
    tenant_id: str | None,
) -> dict[str, dict[str, str]]:
    query = select(TranslationEntry).where(TranslationEntry.tenant_id == tenant_id)
    result = await session.execute(query)
    by_key_locale: dict[str, dict[str, str]] = {}
    for entry in result.scalars().all():
        by_key_locale.setdefault(entry.key, {})[entry.locale] = entry.value
    return by_key_locale


def _merged_value_for_key(
    key: str,
    locale_code: str,
    *,
    master_keys: dict[str, str],
    db_values: dict[str, dict[str, str]],
    static_cache: dict[str, dict[str, str]],
) -> str:
    if locale_code in db_values.get(key, {}):
        return db_values[key][locale_code]
    static_val = static_cache.get(locale_code, {}).get(key)
    if static_val is not None:
        return static_val
    if locale_code == "en":
        return master_keys.get(key, "")
    return ""


async def list_translation_items(
    session: AsyncSession,
    *,
    tenant_id: str | None,
    locale: str | None = None,
    q: str | None = None,
    limit: int = 50,
    offset: int = 0,
) -> dict[str, Any]:
    master_keys = await load_master_keys(session)
    locales = await list_active_locales(session)
    locale_codes = [loc.code for loc in locales] or ["en", "es"]
    if locale:
        locale_codes = [locale.strip().lower()]

    db_values = await _load_db_values(session, tenant_id=tenant_id)
    static_cache = {code: load_static_locale_bundle(code) for code in locale_codes}

    source_keys = sorted(master_keys.keys()) if tenant_id is None else sorted(db_values.keys())
    if tenant_id is None and not source_keys:
        source_keys = sorted(master_keys.keys())

    filtered_keys: list[str] = []
    needle = q.lower().strip() if q else None
    for key in source_keys:
        if not needle:
            filtered_keys.append(key)
            continue
        values_text = master_keys.get(key, "").lower()
        for code in locale_codes:
            values_text += " " + _merged_value_for_key(
                key,
                code,
                master_keys=master_keys,
                db_values=db_values,
                static_cache=static_cache,
            ).lower()
        if needle in key.lower() or needle in values_text:
            filtered_keys.append(key)

    total = len(filtered_keys)
    page_keys = filtered_keys[offset : offset + limit]
    items = []
    for key in page_keys:
        values = {
            code: _merged_value_for_key(
                key,
                code,
                master_keys=master_keys,
                db_values=db_values,
                static_cache=static_cache,
            )
            for code in locale_codes
        }
        items.append(
            {
                "key": key,
                "namespace": key.rsplit(".", 1)[0],
                "values": values,
            }
        )

    return {
        "items": items,
        "total": total,
        "offset": offset,
        "limit": limit,
        "locales": [{"code": loc.code, "name": loc.name, "rtl": loc.rtl} for loc in locales],
    }


async def list_override_items(
    session: AsyncSession,
    tenant_id: str,
    locale: str | None = None,
) -> dict[str, Any]:
    query = select(TranslationEntry).where(TranslationEntry.tenant_id == tenant_id)
    if locale:
        query = query.where(TranslationEntry.locale == locale.strip().lower())
    query = query.order_by(TranslationEntry.key)
    result = await session.execute(query)
    entries = result.scalars().all()
    return {
        "items": [
            {"key": entry.key, "locale": entry.locale, "value": entry.value}
            for entry in entries
        ]
    }


async def upsert_translation(
    session: AsyncSession,
    *,
    tenant_id: str | None,
    locale: str,
    key: str,
    value: str,
    namespace: str | None = None,
    source: str = "inline",
    actor_id: str | None = None,
) -> None:
    locale = locale.strip().lower()
    namespace = namespace or key.rsplit(".", 1)[0]
    result = await session.execute(
        select(TranslationEntry).where(
            TranslationEntry.tenant_id == tenant_id,
            TranslationEntry.locale == locale,
            TranslationEntry.key == key,
        )
    )
    existing = result.scalar_one_or_none()
    if existing:
        existing.value = value
        existing.namespace = namespace
        return

    session.add(
        TranslationEntry(
            id=f"tr-{uuid.uuid4().hex[:12]}",
            tenant_id=tenant_id,
            locale=locale,
            key=key,
            namespace=namespace,
            value=value,
        )
    )


async def register_locale(
    session: AsyncSession,
    *,
    code: str,
    name: str,
    rtl: bool = False,
) -> TranslationLocale:
    code = code.strip().lower()
    existing = await session.execute(select(TranslationLocale).where(TranslationLocale.code == code))
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=409, detail={"error": "locale_exists", "message": "Locale already registered"})
    locale = TranslationLocale(code=code, name=name.strip(), rtl=rtl, active=True)
    session.add(locale)
    return locale


async def get_locale(session: AsyncSession, code: str) -> TranslationLocale:
    code = code.strip().lower()
    result = await session.execute(select(TranslationLocale).where(TranslationLocale.code == code))
    locale = result.scalar_one_or_none()
    if not locale:
        raise HTTPException(status_code=404, detail={"error": "not_found", "message": "Locale not found"})
    return locale


async def update_locale(
    session: AsyncSession,
    *,
    code: str,
    name: str | None = None,
    rtl: bool | None = None,
    active: bool | None = None,
) -> TranslationLocale:
    locale = await get_locale(session, code)
    if name is not None:
        locale.name = name.strip()
    if rtl is not None:
        locale.rtl = rtl
    if active is not None:
        if locale.code in PROTECTED_LOCALE_CODES and not active:
            raise HTTPException(
                status_code=400,
                detail={"error": "protected_locale", "message": "English cannot be deactivated"},
            )
        locale.active = active
    return locale


async def delete_locale(session: AsyncSession, code: str) -> None:
    code = code.strip().lower()
    if code in PROTECTED_LOCALE_CODES:
        raise HTTPException(
            status_code=400,
            detail={"error": "protected_locale", "message": "English cannot be removed"},
        )
    locale = await get_locale(session, code)
    await session.execute(delete(TranslationEntry).where(TranslationEntry.locale == code))
    await session.delete(locale)


async def get_translation_bundle(
    session: AsyncSession,
    locale: str,
    *,
    tenant_id: str | None = None,
) -> dict[str, str]:
    locale = locale.strip().lower()
    static = load_static_locale_bundle(locale)
    result = await session.execute(
        select(TranslationEntry).where(
            TranslationEntry.tenant_id == tenant_id,
            TranslationEntry.locale == locale,
        )
    )
    db_entries = {entry.key: entry.value for entry in result.scalars().all()}
    return {**static, **db_entries}


async def export_translations_xlsx(
    session: AsyncSession,
    *,
    tenant_id: str | None,
    filename: str,
) -> StreamingResponse:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail={"error": "missing_dependency", "message": "openpyxl not installed"},
        ) from exc

    master_keys = await load_master_keys(session)
    locales = await list_active_locales(session)
    locale_codes = [loc.code for loc in locales] or ["en", "es"]
    db_values = await _load_db_values(session, tenant_id=tenant_id)
    static_cache = {code: load_static_locale_bundle(code) for code in locale_codes}

    keys = sorted(master_keys.keys()) if tenant_id is None else sorted(db_values.keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "Translations"
    headers = ["key", "namespace", *locale_codes, "description", "tenant_overridable"]
    ws.append(headers)
    for key in keys:
        row = [key, key.rsplit(".", 1)[0]]
        for code in locale_codes:
            row.append(
                _merged_value_for_key(
                    key,
                    code,
                    master_keys=master_keys,
                    db_values=db_values,
                    static_cache=static_cache,
                )
            )
        row.extend(["", "Y" if tenant_id else "N"])
        ws.append(row)

    if tenant_id is None:
        ws_loc = wb.create_sheet("Locales")
        ws_loc.append(["code", "name", "active"])
        all_locales = await session.execute(select(TranslationLocale).order_by(TranslationLocale.code))
        for loc in all_locales.scalars().all():
            ws_loc.append([loc.code, loc.name, "Y" if loc.active else "N"])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def import_translations_xlsx(
    session: AsyncSession,
    *,
    tenant_id: str | None,
    content: bytes,
    actor_id: str,
) -> int:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail={"error": "missing_dependency", "message": "openpyxl not installed"},
        ) from exc

    wb = load_workbook(BytesIO(content), read_only=True)
    if "Translations" not in wb.sheetnames:
        raise HTTPException(status_code=400, detail={"error": "invalid_file", "message": "Missing Translations sheet"})

    if tenant_id is None and "Locales" in wb.sheetnames:
        ws_loc = wb["Locales"]
        loc_rows = list(ws_loc.iter_rows(values_only=True))
        if len(loc_rows) >= 2:
            loc_headers = [str(h).strip().lower() if h else "" for h in loc_rows[0]]
            code_idx = loc_headers.index("code") if "code" in loc_headers else 0
            name_idx = loc_headers.index("name") if "name" in loc_headers else 1
            active_idx = loc_headers.index("active") if "active" in loc_headers else None
            for row in loc_rows[1:]:
                if not row or not row[code_idx]:
                    continue
                code = str(row[code_idx]).strip().lower()
                name = (
                    str(row[name_idx]).strip()
                    if name_idx is not None and name_idx < len(row) and row[name_idx]
                    else code
                )
                active = True
                if active_idx is not None and active_idx < len(row) and row[active_idx]:
                    active = str(row[active_idx]).strip().upper() in {"Y", "YES", "TRUE", "1"}
                existing = await session.execute(
                    select(TranslationLocale).where(TranslationLocale.code == code)
                )
                loc = existing.scalar_one_or_none()
                if loc:
                    loc.name = name
                    loc.active = active
                else:
                    session.add(TranslationLocale(code=code, name=name, rtl=False, active=active))

    master_keys = await load_master_keys(session)
    ws = wb["Translations"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail={"error": "invalid_file", "message": "No translation rows found"})

    headers = [str(header).strip().lower() if header else "" for header in rows[0]]
    key_idx = headers.index("key") if "key" in headers else 0
    ns_idx = headers.index("namespace") if "namespace" in headers else None
    locale_cols = [
        (index, header)
        for index, header in enumerate(headers)
        if header and header not in {"key", "namespace", "description", "tenant_overridable"}
    ]

    imported = 0
    for row in rows[1:]:
        if not row or not row[key_idx]:
            continue
        key = str(row[key_idx]).strip()
        if tenant_id is None and key not in master_keys:
            continue
        namespace = (
            str(row[ns_idx]).strip()
            if ns_idx is not None and ns_idx < len(row) and row[ns_idx]
            else key.rsplit(".", 1)[0]
        )
        for col_idx, locale in locale_cols:
            if col_idx >= len(row):
                continue
            value = row[col_idx]
            if value is None or str(value).strip() == "":
                continue
            await upsert_translation(
                session,
                tenant_id=tenant_id,
                locale=locale,
                key=key,
                value=str(value),
                namespace=namespace,
                source="import",
                actor_id=actor_id,
            )
            imported += 1
    return imported
